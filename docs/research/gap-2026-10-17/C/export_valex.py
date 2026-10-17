import openpyxl, json, sys
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
# sheet -> (first data row, key column that must be numeric)
SPEC = {
  'P.838-3 Sp.Att': (20, 'G'),
  'P.839-4 Rain_Height': (22, 'C'),
  'P.840-9 A_Clouds': (21, 'F'),
  'P840_9_L': (24, 'E'),
  'P.837-7 Rp': (22, 'E'),
  'P.618-14 A_Rain': (23, 'G'),
  'P.618-14 A_Scint': (21, 'G'),
  'P.618-14 XPD': (21, 'H'),
  'P.618-14 Att_Tot': (26, 'G'),
  'P.676-13 A_Gas_A2_INST': (24, 'F'),
  'P.453-14 Nwet': (21, 'C'),
  'P.836-6 WV': (25, 'G'),
  'P.1511 TOPO': (23, 'C'),
}
out = {}
for name, (r0, keycol) in SPEC.items():
    ws = wb[name]
    rows = []
    for row in ws.iter_rows(min_row=r0, values_only=True):
        d = {}
        for j, v in enumerate(row):
            if v is None: continue
            d[openpyxl.utils.get_column_letter(j + 1)] = v
        kv = d.get(keycol)
        if not isinstance(kv, (int, float)):
            if rows and not d: continue
            if not d: continue
            # stop at first non-numeric row after data started
            if rows: break
            else: continue
        rows.append(d)
    out[name] = rows
    print(name, len(rows), file=sys.stderr)
json.dump(out, open(sys.argv[2], 'w', encoding='utf-8'), default=str)
