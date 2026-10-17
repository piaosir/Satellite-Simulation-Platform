import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1],read_only=True,data_only=True)
ws=wb[sys.argv[2]]
r0=int(sys.argv[3]); r1=int(sys.argv[4])
for i,row in enumerate(ws.iter_rows(min_row=r0,max_row=r1,values_only=True),start=r0):
    cells=[]
    for j,v in enumerate(row):
        if v is None: continue
        if isinstance(v,float): v=('%.10g'%v)
        cells.append('%s=%s'%(openpyxl.utils.get_column_letter(j+1),str(v)[:40]))
    if cells: print(i,' | '.join(cells)[:900])
